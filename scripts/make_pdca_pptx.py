#!/usr/bin/env python3
"""生成BF2+BFM4 PDCA幻灯片PPTX文件"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

# 颜色定义
DARK = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x3B, 0x82, 0xF6)
GREEN = RGBColor(0x22, 0xC5, 0x5E)
YELLOW = RGBColor(0xEA, 0xB3, 0x08)
PURPLE = RGBColor(0xA8, 0x55, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF8, 0xFA, 0xFC)
GRAY = RGBColor(0x64, 0x74, 0x8B)
DARK_TEXT = RGBColor(0x1E, 0x29, 0x3B)

# ─── 顶部标题栏 ───
bg = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(1.0))  # 1=rectangle
bg.fill.solid()
bg.fill.fore_color.rgb = DARK
bg.line.fill.background()

txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(10), Inches(0.7))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "BF2+BFM4 4-Wire Process – PDCA Summary"
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = WHITE

p2 = tf.add_paragraph()
p2.text = "PDCA Cycle · Process Improvement"
p2.font.size = Pt(14)
p2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

# PDCA四块内容
cards = [
    {
        "letter": "P", "phase": "Plan · Cost Advantage & Risk", "tag": "Strategy",
        "color": BLUE, "bg_light": RGBColor(0xEF, 0xF6, 0xFF),
        "border": RGBColor(0xBF, 0xDB, 0xFE),
        "items": [
            ("187–240 RMB/ton cost saving, payback <1 year", True),
            ("Process replacement: BF2+BFM4 vs BFM9/BFM6", False),
            ("⚠ Risk: 4-wire process may increase defects (core inversion, particle/dirt issues)", True),
        ]
    },
    {
        "letter": "D", "phase": "Do · Trial & Improvement", "tag": "Execution",
        "color": GREEN, "bg_light": RGBColor(0xF0, 0xFD, 0xF4),
        "border": RGBColor(0xBB, 0xF7, 0xD0),
        "items": [
            ("Pilot + plant trial validation (Bardec & CBSC)", False),
            ("AWC 2.0 applied on WWD machines", True),
            ("Improved winding layout & reduced length difference", False),
        ]
    },
    {
        "letter": "C", "phase": "Check · Validation Results", "tag": "Results",
        "color": YELLOW, "bg_light": RGBColor(0xFE, 0xFC, 0xE8),
        "border": RGBColor(0xFD, 0xE6, 0x8A),
        "items": [
            ("Trial lot quality: HASP, SLT, formation ✓ OK", False),
            ("Stable process with low defect rate", True),
            ("Customer performance validated OK", False),
        ]
    },
    {
        "letter": "A", "phase": "Act · Scale-up Plan", "tag": "Next Steps",
        "color": PURPLE, "bg_light": RGBColor(0xFA, 0xF5, 0xFF),
        "border": RGBColor(0xE9, 0xD5, 0xFF),
        "items": [
            ("Scale up application of BF2+BFM4 process", False),
            ("Expand capacity to 24 production lines", True),
            ("Start ramp-up beginning of July", False),
            ("Implement stepwise ramp-up plan (capacity & yield increase)", False),
        ]
    }
]

positions = [
    (Inches(0.4), Inches(1.2)),
    (Inches(6.7), Inches(1.2)),
    (Inches(0.4), Inches(4.35)),
    (Inches(6.7), Inches(4.35)),
]

for i, card in enumerate(cards):
    left, top = positions[i]
    w, h = Inches(6.2), Inches(3.0)
    
    # Card background
    rect = slide.shapes.add_shape(1, left, top, w, h)
    rect.fill.solid()
    rect.fill.fore_color.rgb = card["bg_light"]
    rect.line.color.rgb = card["border"]
    rect.line.width = Pt(1)
    
    # Letter + Phase
    txBox = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), Inches(0.5), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = card["letter"]
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.fill.solid()
    # letter background via second shape
    letter_bg = slide.shapes.add_shape(1, left + Inches(0.25), top + Inches(0.2), Inches(0.45), Inches(0.45))
    letter_bg.fill.solid()
    letter_bg.fill.fore_color.rgb = card["color"]
    letter_bg.line.fill.background()
    
    # Phase text
    txBox = slide.shapes.add_textbox(left + Inches(0.9), top + Inches(0.22), Inches(4.5), Inches(0.45))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = card["phase"]
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = card["color"]
    
    # Tag
    tag = slide.shapes.add_textbox(left + Inches(5.0), top + Inches(0.2), Inches(1.0), Inches(0.35))
    tf = tag.text_frame
    p = tf.paragraphs[0]
    p.text = card["tag"]
    p.font.size = Pt(10)
    p.font.color.rgb = card["color"]
    p.alignment = PP_ALIGN.RIGHT
    
    # Items
    y_offset = Inches(0.85)
    for item_text, is_highlight in card["items"]:
        txBox = slide.shapes.add_textbox(left + Inches(0.35), top + y_offset, Inches(5.5), Inches(0.4))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = item_text
        p.font.size = Pt(13)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)
        
        if is_highlight:
            p.font.bold = True
        
        y_offset += Inches(0.45)

# ─── 底部条 ───
footer_bg = slide.shapes.add_shape(1, Inches(0), Inches(7.1), Inches(13.333), Inches(0.4))
footer_bg.fill.solid()
footer_bg.fill.fore_color.rgb = RGBColor(0xF1, 0xF5, 0xF9)
footer_bg.line.fill.background()

txBox = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(6), Inches(0.4))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "BF2+BFM4 4-Wire Process · PDCA Summary"
p.font.size = Pt(11)
p.font.color.rgb = GRAY

txBox = slide.shapes.add_textbox(Inches(8.5), Inches(7.1), Inches(4.5), Inches(0.4))
tf = txBox.text_frame
p = tf.paragraphs[0]
p.text = "P · D · C · A          Continuous Improvement"
p.font.size = Pt(11)
p.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
p.alignment = PP_ALIGN.RIGHT

# 保存
output_path = os.path.expanduser("~/.hermes/reports/BF2_BFM4_PDCA_Summary.pptx")
prs.save(output_path)
print(f"✅ PPTX已保存: {output_path}")
print(f"文件大小: {os.path.getsize(output_path)/1024:.0f} KB")
